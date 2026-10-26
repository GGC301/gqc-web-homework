# -*- coding: utf-8 -*-
"""图像解决方案实验报告生成脚本
复刻菜鸟教程官网手机版：iconfont 字体图标 + 内联 SVG 实现自适应宽度固定底部下导航
"""
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import os

doc = Document()

# ============ 封面信息 ============
title = doc.add_heading('《高级网页设计I》图像解决方案 项目报告', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

info = doc.add_paragraph()
info.add_run('学  号：').bold = True
info.add_run('24001010301\n')
info.add_run('姓  名：').bold = True
info.add_run('关泉才\n')
info.add_run('班  级：').bold = True
info.add_run('24001班')

# ============ 项目选题 ============
doc.add_heading('项目选题名称', level=1)
doc.add_paragraph('参考菜鸟教程官网（www.runoob.com），复刻其手机版页面。重点实现网站的底部下导航：'
                  '使用 iconfont 字体图标 + 内联 SVG 作为图像解决方案，'
                  '通过 CSS 实现自适应宽度的固定底部导航栏（5 个导航项：首页 / 教程 / 论坛 / 收藏 / 我的）。')

# ============ 实施过程 ============
doc.add_heading('实施过程 - 项目需求与实现', level=1)
doc.add_paragraph(
    '根据菜鸟教程官网手机版设计稿，使用 HTML 语义化标签和 CSS Flex 布局完成页面复刻。'
    '图像解决方案采用两种技术：'
    '① iconfont 字体图标（从 iconfont.cn 在线项目 gqc301 引入，类名 icon- 前缀，使用 @font-face 嵌入 woff2/ttf 字体文件）；'
    '② 内联 SVG（搜索放大镜图标、banner 渐变背景）。'
    '下导航使用 position:fixed 固定在页面底部，display:flex + flex:1 实现 5 项等宽自适应；'
    '首页为激活态（绿色 + 图标放大），其余为灰色。')

shots = r'C:\Users\GGC\Desktop\page\images-solution\shots'
orig = r'C:\Users\GGC\Desktop\page\web-elements\shots'

# ============ 原图与实现图 ============
doc.add_heading('原图（菜鸟教程官网手机版）', level=2)
doc.add_paragraph('菜鸟教程官网首页（www.runoob.com）：')
doc.add_picture(os.path.join(orig, 'orig-home.png'), width=Inches(6))

doc.add_heading('实现图（菜鸟教程官网手机版复刻）', level=2)
doc.add_paragraph('完整页面实现图（顶部 logo+搜索、SVG 渐变 banner、iconfont 快捷入口、教程列表、固定底部下导航）：')
doc.add_picture(os.path.join(shots, 'impl-full.png'), width=Inches(6))

doc.add_heading('下导航放大图（自适应宽度 + 固定底部）', level=2)
doc.add_paragraph('底部下导航 5 项：首页 / 教程 / 论坛 / 收藏 / 我的，图标为 iconfont 字体图标，'
                  '首页为激活态（绿色 #73a857 + 图标放大）：')
doc.add_picture(os.path.join(shots, 'impl-nav.png'), width=Inches(6))

# ============ 核心代码 ============
doc.add_heading('项目核心代码', level=1)
doc.add_paragraph(
    '代码中类名统一使用 gqc- 前缀（guanquancai 缩写），包含学号 24001010301 和姓名信息，'
    '使用 class 选择器设计全局 CSS。字体图标从 iconfont 项目引入（@font-face 字体文件 + icon- 类名）。')

doc.add_heading('1. iconfont 字体图标引入（head 中）', level=2)
doc.add_paragraph('''<link rel="stylesheet" href="font/iconfont.css">
<!-- iconfont.css 内容（由 iconfont.cn 自动生成） -->
@font-face {
  font-family: "iconfont";
  src: url('iconfont.woff2?t=...') format('woff2'),
       url('iconfont.woff?t=...') format('woff'),
       url('iconfont.ttf?t=...') format('truetype');
}
.iconfont { font-family: "iconfont" !important; font-style: normal;
  -webkit-font-smoothing: antialiased; }''')

doc.add_heading('2. 下导航 HTML 结构（5 项 iconfont 字体图标）', level=2)
doc.add_paragraph('''<nav class="gqc-m-tabbar">
  <a href="#" class="gqc-active">
    <i class="iconfont icon-shouye"></i><span>首页</span>
  </a>
  <a href="#">
    <i class="iconfont icon-gengduo"></i><span>教程</span>
  </a>
  <a href="#">
    <i class="iconfont icon-luntan"></i><span>论坛</span>
  </a>
  <a href="#">
    <i class="iconfont icon-xihuan"></i><span>收藏</span>
  </a>
  <a href="#">
    <i class="iconfont icon-gerenzhongxin-zhihui"></i><span>我的</span>
  </a>
</nav>''')

doc.add_heading('3. 下导航 CSS（自适应宽度 + 固定底部 + 字体图标修饰）', level=2)
doc.add_paragraph('''/* 固定底部 + flex 自适应等宽 */
.gqc-m-tabbar { position:fixed; bottom:0; left:0; right:0; height:56px;
  background:#fff; border-top:1px solid #e8e8e8; display:flex; z-index:100; }
.gqc-m-tabbar a { flex:1; display:flex; flex-direction:column;
  align-items:center; justify-content:center; gap:3px; color:#999; min-width:0; }
/* 字体图标默认样式（CSS 修饰：大小、颜色） */
.gqc-m-tabbar .iconfont { font-size:22px; line-height:1; }
.gqc-m-tabbar span { font-size:10px; }
/* 激活态：绿色 + 图标放大（CSS 修饰字体图标） */
.gqc-m-tabbar a.gqc-active { color:#73a857; }
.gqc-m-tabbar a.gqc-active .iconfont { font-size:26px; }''')

doc.add_heading('4. 内联 SVG（搜索图标 + banner 渐变背景）', level=2)
doc.add_paragraph('''<!-- 搜索放大镜 SVG -->
<svg viewBox="0 0 24 24"><path d="M15.5 14h-.79l-.28-.27a6.5 6.5 0 1 0-.7.7l.27.28v.79l5 4.99L20.49 19l-4.99-5zm-6 0A4.5 4.5 0 1 1 14 9.5 4.5 4.5 0 0 1 9.5 14z"/></svg>

<!-- banner 线性渐变 SVG -->
<svg viewBox="0 0 400 120" preserveAspectRatio="none">
  <defs>
    <linearGradient id="gqcGrad" x1="0" y1="0" x2="1" y2="1">
      <stop offset="0%" stop-color="#8bc34a"/>
      <stop offset="100%" stop-color="#3f7d20"/>
    </linearGradient>
  </defs>
  <rect width="400" height="120" fill="url(#gqcGrad)"/>
</svg>''')

# ============ 测试与发布 ============
doc.add_heading('测试与发布', level=1)
doc.add_paragraph('在 Microsoft Edge 浏览器中打开 index.html，页面完整渲染：'
                  '顶部 logo 与搜索框、SVG 渐变 banner、iconfont 快捷入口、热门教程列表、'
                  '底部固定下导航 5 项均正常显示。滚动页面后下导航仍固定在底部（position:fixed 生效），'
                  '首页激活态为绿色。')
doc.add_paragraph('页面测试图（下导航固定显示在底部）：')
doc.add_picture(os.path.join(shots, 'impl-nav.png'), width=Inches(6))

# ============ 项目总结 ============
doc.add_heading('项目总结', level=1)
doc.add_paragraph(
    '本次实验通过复刻菜鸟教程官网手机版，实现了图像的两种解决方案：'
    '①iconfont 字体图标——从 iconfont.cn 获取在线图标项目（gqc301，20 个图标），'
    '通过 @font-face 引入 woff2/ttf 字体文件，用 <i class="iconfont icon-xxx"></i> 方式在页面中使用；'
    '②内联 SVG——搜索放大镜图标与 banner 线性渐变背景直接用 SVG 代码绘制，无需图片文件。'
    '下导航使用 position:fixed + display:flex + flex:1 实现自适应宽度、固定底部的效果，'
    '并用 CSS 对字体图标进行修饰（颜色、大小、激活态）。'
    '主要收获：①掌握了字体图标的获取与引入流程；②理解 SVG 内联绘制与渐变定义；'
    '③学会了 flex 实现自适应等宽导航栏；④巩固了 fixed 定位实现固定底栏。'
    '使用的 AI 方法：使用豆包 AI 辅助生成 HTML 结构和 CSS 样式，参考菜鸟教程官网设计稿进行人工调整。'
    '学号 24001010301 关泉才。'
)

out = r'C:\Users\GGC\Downloads\24001010301-关泉才-图像解决方案.docx'
doc.save(out)
print('saved', out)
