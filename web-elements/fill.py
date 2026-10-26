from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
from copy import deepcopy
import os

doc = Document(r'C:\Users\GGC\Desktop\page\web-elements\template.docx')
shots = r'C:\Users\GGC\Desktop\page\web-elements\shots'

# 1. 填表格
t = doc.tables[0]
# 行0: 学号 姓名
cell0 = t.rows[0].cells[0]
cell0.text = '  学  号：24001010301          姓   名：关泉才'
# 行1: 班级
t.rows[1].cells[1].text = '24001班'
# 行2: 项目选题名称
t.rows[2].cells[1].text = '参考菜鸟教程官网，完成顶部导航栏、课程卡片列表、左侧目录+右侧内容区三个网页元素设计'

# 2. 在指定段落后插入图片的辅助函数
def add_image_after(paragraph, img_path, width_in=5.5):
    new_p = paragraph.insert_paragraph_before('')
    # insert_paragraph_before 是在前面插，我们需要在后面
    # 改用：在目标段落后面加新段落
    pass

# 直接遍历段落，找到标记段，在它后面插图
paras = doc.paragraphs

# 找到各标记段索引
def find_idx(text_start):
    for i, p in enumerate(paras):
        if p.text.strip().startswith(text_start):
            return i
    return -1

# 图片映射：标记文本 -> 图片文件
img_map = [
    ('网页元素1原图', 'orig-home.png'),
    ('网页元素1实现图', 'impl-full.png'),
    ('网页元素2原图', 'orig-home.png'),
    ('网页元素2实现图', 'impl-full.png'),
    ('网页元素3原图', 'orig-home.png'),
    ('网页元素3实现图', 'impl-full.png'),
    ('网页元素一测试图', 'impl-full.png'),
]

# 从后往前插，避免索引变化
for label, img_file in reversed(img_map):
    idx = find_idx(label)
    if idx < 0:
        print('NOT FOUND', label)
        continue
    p = paras[idx]
    # 在该段落后面插入新段落放图片
    new_p = deepcopy(p._p)
    # 清空内容
    for child in list(new_p):
        new_p.remove(child)
    p._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    np = Paragraph(new_p, p._parent)
    run = np.add_run()
    run.add_picture(os.path.join(shots, img_file), width=Inches(5.5))
    print('inserted', label, img_file)

# 3. 填代码
code_map = {
    '网页元素1核心代码': '''<nav class="gqc-navbar">
  <div class="gqc-navbar__logo">
    <div class="gqc-navbar__logo-mark">run</div>
    <span>菜鸟教程</span>
  </div>
  <div class="gqc-navbar__menu">
    <a href="#">首页</a><a href="#">HTML教程</a><a href="#">CSS教程</a>
    <a href="#">JavaScript</a><a href="#">Python</a><a href="#">SQL</a>
  </div>
  <div class="gqc-navbar__search">
    <input type="text" placeholder="输入教程关键字"><button>搜索</button>
  </div>
</nav>

.gqc-navbar { background:#222f3e; display:flex; justify-content:space-between;
  align-items:center; padding:0 40px; height:60px; color:#fff; }
.gqc-navbar__menu { display:flex; gap:28px; }''',

    '网页元素2核心代码': '''<div class="gqc-card-list">
  <div class="gqc-card">
    <div class="gqc-card__thumb gqc-card__thumb--green">HTML</div>
    <div class="gqc-card__body">
      <div class="gqc-card__title">HTML 教程</div>
      <div class="gqc-card__desc">超文本标记语言...</div>
    </div>
  </div>
</div>

.gqc-card-list { display:grid; grid-template-columns:repeat(4,1fr); gap:20px; }
.gqc-card { background:#fff; border-radius:8px; overflow:hidden;
  box-shadow:0 2px 8px rgba(0,0,0,.08); transition:transform .2s; }
.gqc-card:hover { transform:translateY(-4px); }''',

    '网页元素3核心代码': '''<section class="gqc-layout">
  <aside class="gqc-sidebar">
    <div class="gqc-sidebar__item gqc-sidebar__item--active">HTML 简介</div>
    <div class="gqc-sidebar__item">HTML 编辑器</div>
    <div class="gqc-sidebar__item">HTML 基础</div>
  </aside>
  <div class="gqc-content">
    <h2>HTML 简介</h2>
    <p>关泉才 24001010301 - 高级网页设计I 实验作业</p>
  </div>
</section>

.gqc-layout { display:flex; gap:24px; max-width:1200px; margin:40px auto; }
.gqc-sidebar { flex:0 0 220px; background:#fff; border-radius:8px; }
.gqc-content { flex:1; background:#fff; border-radius:8px; padding:24px; }''',
}

# 重新获取段落列表
paras = doc.paragraphs
for label, code in code_map.items():
    idx = -1
    for i, p in enumerate(paras):
        if p.text.strip().startswith(label):
            idx = i
            break
    if idx < 0:
        print('CODE NOT FOUND', label)
        continue
    p = paras[idx]
    # 在后面插代码段落
    new_p = deepcopy(p._p)
    for child in list(new_p):
        new_p.remove(child)
    p._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    np = Paragraph(new_p, p._parent)
    run = np.add_run(code)
    run.font.size = Pt(9)
    print('code inserted', label)

# 4. 项目总结：在"项目总结"标题后的说明段落后面加内容
for i, p in enumerate(doc.paragraphs):
    if '提交项目报告者' in p.text:
        p.add_run('\n\n本次实验通过复刻菜鸟教程官网的三个典型网页元素，巩固了HTML语义化标签和CSS Flex/Grid布局。'
                  '主要收获：①掌握Flex导航栏的两端对齐；②用Grid做卡片网格；③BEM命名和class复用。'
                  '使用豆包AI辅助生成代码，人工对照官网调整样式。学号24001010301 关泉才。')
        break

out = r'C:\Users\GGC\Downloads\24001010301-关泉才-网页元素设计.docx'
doc.save(out)
print('saved', out)
